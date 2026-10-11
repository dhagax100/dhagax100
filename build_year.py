#!/usr/bin/env python3
"""Runs all 12 months of ONE year, building the expensive full-history
Daily engine + control-checkpoint timeline exactly ONCE (shared across all
12 months) instead of 12 times -- this is the item #1 speed fix. Produces,
for THAT YEAR ONLY (never combined across years, per instruction):
  - YEAR_report.xlsx   -- one sheet per month, locked 4-column table style
  - YEAR_trades.csv    -- ONE comprehensive row per real trade: date,
    timeframe, side, control state at impact, entry hour (Riyadh), SL
    size (pips), result, R-multiple, and the full structural_notes text
    (which includes item #10's diagnostic comparison note whenever it
    fired, item #11's prior-4H note, item #25's instant-stop note, etc.)
  - YEAR_month_01.md .. YEAR_month_12.md -- the raw markdown per month

Usage: python build_year.py YEAR
Example: python build_year.py 2015
"""
import sys, csv, time
from pathlib import Path
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo
from types import SimpleNamespace

_REPO_ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(_REPO_ROOT / "reference_daily"))
sys.path.insert(0, str(_REPO_ROOT))

import verify_day as vd
import all_tf_combined_generator as atc
import weekly_ob_generator as wob
import weekly_combined_generator as wc
import daily_combined_generator as dc
import build_report as br

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

RYD = ZoneInfo("Asia/Riyadh")
UTC = ZoneInfo("UTC")
CLOSE_TZ = ZoneInfo("America/New_York")


def build_shared(csv_path, final_cutoff_str):
    t0 = time.time()
    y, m, d = (int(x) for x in final_cutoff_str.split("-"))
    cutoff = datetime(y, m, d, 23, 59, 59, tzinfo=RYD).astimezone(UTC)

    full_minutes, _ = wob.load_minutes(csv_path, UTC, "bid")
    full_minutes, _ = atc.strip_pre_week_open(full_minutes, RYD)
    full_minutes = [x for x in full_minutes if x.t < cutoff]
    full_mt = [x.t for x in full_minutes]
    print(f"[{time.time()-t0:.1f}s] loaded full history: {len(full_minutes)} minutes", file=sys.stderr)

    bars_d = dc.aggregate_days(full_minutes, CLOSE_TZ, 17, RYD)
    d_engine = wc.WeeklyCombinedEngine(full_minutes, bars_d, origin_gap_window=None)
    d_engine.run()
    atc.mark_open_inside_trigger(d_engine)
    atc.mark_superseded_same_leg(d_engine, full_minutes, full_mt)
    d_zones_full = SimpleNamespace(
        ob_zones=list(d_engine.ob_zones), rb_zones=list(d_engine.rb_zones),
        fvg_zones=list(d_engine.fvg_zones), vi_zones=list(d_engine.vi_zones),
        w=d_engine.w, events=d_engine.events, msses=list(d_engine.msses))
    print(f"[{time.time()-t0:.1f}s] Daily engine built (full history, shared across all 12 months)", file=sys.stderr)

    atc.daily_controlling_bias(d_zones_full, cutoff, full_minutes, full_mt)
    gate_control_checkpoints = atc.compute_control_timeline(d_zones_full, full_minutes, full_mt, RYD, cutoff)
    print(f"[{time.time()-t0:.1f}s] control checkpoints computed", file=sys.stderr)

    return SimpleNamespace(t0=t0, cutoff=cutoff, full_minutes=full_minutes, full_mt=full_mt,
                            d_engine=d_engine, d_zones_full=d_zones_full,
                            gate_control_checkpoints=gate_control_checkpoints)


def build_month_ctx(shared, year, month, lead_months=2):
    t0 = shared.t0
    short_start = (datetime(year, month, 1, tzinfo=RYD) - timedelta(days=30 * lead_months)).astimezone(UTC)
    short_minutes = [x for x in shared.full_minutes if x.t >= short_start]
    short_mt = [x.t for x in short_minutes]

    bars_h4 = dc.aggregate_hours(short_minutes, 4, CLOSE_TZ, 17, RYD)
    h4_engine = wc.WeeklyCombinedEngine(short_minutes, bars_h4, origin_gap_window=None)
    h4_engine.run()
    atc.mark_open_inside_trigger(h4_engine)
    atc.mark_superseded_same_leg(h4_engine, short_minutes, short_mt)
    atc.apply_daily_bias_gate(h4_engine, "h4", shared.d_zones_full, RYD, short_minutes, short_mt,
                               control_checkpoints=shared.gate_control_checkpoints)
    atc.exclude_old_intraday_zones(h4_engine)
    tf_full_h4 = SimpleNamespace(ob_zones=list(h4_engine.ob_zones), rb_zones=list(h4_engine.rb_zones),
                                  fvg_zones=list(h4_engine.fvg_zones), vi_zones=list(h4_engine.vi_zones),
                                  events=h4_engine.events, w=h4_engine.w,
                                  ambiguous_tie_bars=set(h4_engine.ambiguous_tie_bars))

    bars_h1 = dc.aggregate_hours(short_minutes, 1, CLOSE_TZ, 17, RYD)
    h1_engine = wc.WeeklyCombinedEngine(short_minutes, bars_h1, origin_gap_window=None)
    h1_engine.run()
    atc.mark_open_inside_trigger(h1_engine)
    atc.mark_superseded_same_leg(h1_engine, short_minutes, short_mt)
    atc.apply_daily_bias_gate(h1_engine, "h1", shared.d_zones_full, RYD, short_minutes, short_mt,
                               control_checkpoints=shared.gate_control_checkpoints)
    atc.exclude_old_intraday_zones(h1_engine)
    tf_full_h1 = SimpleNamespace(ob_zones=list(h1_engine.ob_zones), rb_zones=list(h1_engine.rb_zones),
                                  fvg_zones=list(h1_engine.fvg_zones), vi_zones=list(h1_engine.vi_zones),
                                  events=h1_engine.events, w=h1_engine.w,
                                  ambiguous_tie_bars=set(h1_engine.ambiguous_tie_bars))

    bars5 = dc.aggregate_minutes(short_minutes, 5)
    e5 = wc.WeeklyCombinedEngine(short_minutes, bars5, origin_gap_window=None)
    e5.run()

    return SimpleNamespace(t0=t0, cutoff=shared.cutoff, full_minutes=shared.full_minutes, full_mt=shared.full_mt,
                            d_engine=shared.d_engine, d_zones_full=shared.d_zones_full,
                            gate_control_checkpoints=shared.gate_control_checkpoints,
                            short_minutes=short_minutes, short_mt=short_mt,
                            h4_engine=h4_engine, h1_engine=h1_engine, e5=e5,
                            tf_full_h4=tf_full_h4, tf_full_h1=tf_full_h1)


def main():
    if len(sys.argv) < 2:
        print("Usage: python build_year.py YEAR")
        sys.exit(1)
    YEAR = int(sys.argv[1])

    # final cutoff: 10 days past Dec 31, so December's own trailing-open
    # trades can resolve (same as item #22's fix).
    final_cutoff = (date(YEAR, 12, 31) + timedelta(days=10)).isoformat()
    shared = build_shared(vd.DEFAULT_CSV, final_cutoff)

    wb = Workbook()
    wb.remove(wb.active)

    grand_total_trades = grand_tp = grand_sl = 0
    grand_net_r = 0.0
    all_trade_rows = []

    for month in range(1, 13):
        mt0 = time.time()
        print(f"\n=== Building {br.MONTH_NAMES[month]} {YEAR} ===", file=sys.stderr)
        mctx = build_month_ctx(shared, YEAR, month, lead_months=2)
        out_path = f"{YEAR}_month_{month:02d}.md"
        total_trades, total_tp, total_sl, net_r, entered_rows = br.build_month_from_ctx(
            mctx, YEAR, month, out_path=out_path
        )
        print(f"[{time.time()-shared.t0:.1f}s] month {month} done ({time.time()-mt0:.1f}s this month)", file=sys.stderr)
        grand_total_trades += total_trades
        grand_tp += total_tp
        grand_sl += total_sl
        grand_net_r += net_r

        for r in entered_rows:
            entry_riyadh = r.get("entry_riyadh") or ""
            entry_hour = entry_riyadh.split(" ")[1][:5] if " " in entry_riyadh else ""
            all_trade_rows.append({
                "date": r.get("date", ""),
                "poi": r.get("poi", ""),
                "timeframe": r.get("tf", ""),
                "side": r.get("side", ""),
                "control_state_at_impact": r.get("control_state_at_impact", ""),
                "entry_time_riyadh": entry_riyadh,
                "entry_hour_riyadh": entry_hour,
                "entry_price": r.get("entry_price", ""),
                "sl_price": r.get("sl_price", ""),
                "tp_price": r.get("tp_price", ""),
                "sl_pips": r.get("r_pips", ""),
                "mfe_pips": r.get("mfe_pips", ""),
                "mae_pips": r.get("mae_pips", ""),
                "result": r.get("result", ""),
                "r_multiple": r.get("r_multiple", ""),
                "exit_riyadh": r.get("exit_riyadh", ""),
                "mechanism": r.get("mechanism", ""),
                "structural_notes": r.get("structural_notes", ""),
            })

        with open(out_path, "r", encoding="utf-8") as f:
            text = f.read()
        rows = []
        for line in text.splitlines():
            if line.startswith("| ") and not line.startswith("| Date") and not line.startswith("|---"):
                parts = [p.strip() for p in line.strip().strip("|").split("|")]
                if len(parts) == 4:
                    rows.append(parts)
        ws = wb.create_sheet(title=br.MONTH_NAMES[month])
        ws.append(["Date", "Bias", "Setups", "Entries"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        for col, width in zip("ABCD", [10, 60, 60, 90]):
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    summary = wb.create_sheet(title="Summary", index=0)
    summary.append(["Metric", "Value"])
    summary.append(["Year", YEAR])
    summary.append(["Total trades", grand_total_trades])
    summary.append(["TP", grand_tp])
    summary.append(["SL", grand_sl])
    summary.append(["Net R", grand_net_r])
    summary.column_dimensions["A"].width = 20
    summary.column_dimensions["B"].width = 15

    xlsx_name = f"{YEAR}_report.xlsx"
    wb.save(xlsx_name)

    csv_name = f"{YEAR}_trades.csv"
    fieldnames = ["date", "poi", "timeframe", "side", "control_state_at_impact",
                  "entry_time_riyadh", "entry_hour_riyadh", "entry_price", "sl_price",
                  "tp_price", "sl_pips", "mfe_pips", "mae_pips", "result", "r_multiple",
                  "exit_riyadh", "mechanism", "structural_notes"]
    with open(csv_name, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(all_trade_rows)

    print(f"\n\nDONE -- saved {xlsx_name} and {csv_name}", file=sys.stderr)
    print(f"{YEAR}: {grand_total_trades} trades, {grand_tp} TP, {grand_sl} SL, net {grand_net_r:+.0f}R", file=sys.stderr)


if __name__ == "__main__":
    main()
